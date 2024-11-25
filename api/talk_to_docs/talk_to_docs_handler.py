import os
from io import BytesIO
import logging

import faiss
import numpy as np

from azure.storage.blob import BlobServiceClient
from azure.core.credentials import AzureSasCredential

from gpt4all import GPT4All
from nomic import embed

from pypdf import PdfReader
from docx import Document
from odf.opendocument import load
from odf.text import P
import openpyxl

logging.basicConfig(level=logging.INFO, format='%(asctime)s - %(levelname)s - %(message)s')

logger = logging.getLogger("azure.core.pipeline.policies.http_logging_policy")
logger.setLevel(logging.WARNING)


class FAISSConnector():
    def __init__(self, dimensionality):
        self.index = faiss.IndexFlatL2(dimensionality)
        self.next_index = 0 # Keep track of indices for vector entries

        # store 
        self.data_dict = {"names": [],
                          "chunks": [],
                          "indices": []}

    async def insert_data(self, data):
        logging.info("Inserting embeddings...")
        embeddings = np.array(data['embeddings']).astype('float32')
        self.index.add(embeddings)

        # Add text, index data, and document name to the data_dict dictionary
        for i in range(len(data['embeddings'])):
            self.data_dict['names'].append(data['name'])
            self.data_dict['chunks'].append(data['chunks'][i])
            self.data_dict['indices'].append(self.next_index)
            self.next_index += 1

    async def search(self, query_embedding, n_results):
        logging.info(f"Performing vector search...")
        query_vector = np.array(query_embedding).reshape(1, -1).astype('float32')
        distances, indices = self.index.search(query_vector, n_results)
        return distances, indices


async def get_envs():
    # Get LM & EM model variables
    embedding_model = os.getenv('EmbeddingModel', 'nomic-embed-text-v1.5')
    language_model = os.getenv('LanguageModel', 'Phi-3-mini-4k-instruct.Q4_0.gguf')

    # Get model storage account variables
    model_storage_account = os.getenv('LLMStorageAccName')
    model_sas_token = os.getenv('LLMStorageSasToken')
    model_container_name = os.getenv('LLMStorageContainerName')

    if not all([embedding_model, language_model, model_storage_account, model_sas_token, model_container_name]):
        logging.error("Some essential environment variables are missing. Please check 'LLMStorage' variables.")
        return False
    
    # Get document storage account variables
    docs_storage_account = os.getenv('AzureStorageAccName', '')
    docs_sas_token = os.getenv('AzureStorageSasToken', '')
    docs_container_name = os.getenv('AzureBlobContainerName', '')

    if not all([docs_storage_account, docs_sas_token, docs_container_name]):
        logging.error("Some essential environment variables are missing. Please check 'AzureStorage' variables.")
        return False
    
    # Get LM setting variables
    ctx_window = int(os.getenv('ContextWindow', '2048'))
    lm_context = os.getenv('LanguageModelContext', '')
    max_tokens = int(os.getenv('MaxTokens', '200'))
    temp = float(os.getenv('LMTemperature', '0.7'))
    top_k = int(os.getenv('TopK', '40'))
    top_p = float(os.getenv('TopP', '0.4'))
    min_p = float(os.getenv('MinP', '0.0'))
    repeat_penalty = float(os.getenv('RepeatPenalty', '1.18'))
    repeat_last_n = int(os.getenv('RepeatLastN', '64'))
    n_batch = int(os.getenv('NBatch', '8'))
    n_predict = int(os.getenv('NPredict', '0'))
    streaming = os.getenv('Streaming', 'False').lower() == 'true'
    if n_predict == 0:
        n_predict = None

    if not lm_context:
        logging.warning("No language model context given.")

    envs = {
        'models': {
            'embedding_model': embedding_model,
            'language_model': language_model
        },
        'models_storage': {
            'model_storage_account': model_storage_account,
            'model_sas_token': model_sas_token,
            'model_container_name': model_container_name
        },
        'doc_storage': {
            'docs_storage_account': docs_storage_account,
            'docs_sas_token': docs_sas_token,
            'docs_container_name': docs_container_name
        },
        'model_settings': {
            'ctx_window': ctx_window,
            'lm_context': lm_context,
            'max_tokens': max_tokens,
            'temp': temp,
            'top_k': top_k,
            'top_p': top_p,
            'min_p': min_p,
            'repeat_penalty': repeat_penalty,
            'repeat_last_n': repeat_last_n,
            'n_batch': n_batch,
            'n_predict': n_predict,
            'streaming': streaming
        }
    }

    return envs


async def parse_pdf(document):
    reader = PdfReader(document)
    full_text = ''
    for page in reader.pages:
        full_text += page.extract_text()
    return full_text


async def parse_docx(document):
    doc = Document(document)
    full_text = ''
    for paragraph in doc.paragraphs:
        full_text += paragraph.text
    return full_text


async def parse_odt(document): 
    doc = load(document)
    full_text = ''
    for paragraph in doc.getElementsByType(P):
        full_text += paragraph.firstChild.data
    return full_text


async def parse_txt_csv(document):
    document.seek(0)
    full_text = document.read().decode('utf-8')
    return full_text


async def parse_xlsx(document):
    wb = openpyxl.load_workbook(document)
    sheet = wb.active
    full_text = ""
    for row in sheet.iter_rows(values_only=True):
        full_text += ' '.join(map(str, row)) + '\n'
    return full_text


async def process_document(file_data, suffix, name=''):
    """ Function to gather and process relevant data from document """
    # Parse the file, extract text and metadata
    logging.info(f"Parsing document...")
    try:
        # Using FileStorage object to process file from memory instead of storage
        if suffix == ".pdf":
            text = await parse_pdf(file_data) 
        elif suffix == ".docx":
            text = await parse_docx(file_data)
        elif suffix == ".odt":
            text = await parse_odt(file_data)
        elif suffix in {".txt", ".csv"}:
            text = await parse_txt_csv(file_data)
        elif suffix == ".xlsx":
            text = await parse_xlsx(file_data)
        else:
            raise Exception
    except:
        logging.warning("Document may be incompatible or corrupt, skipping...")
        return
    finally:
        file_data.close()

    if not name:
        name = file_data.filename  # Document name

    # Separate text into list of words
    words = text.split()

    # Create text chunks
    text_chunks = []
    chunk = []
    counter = 0
    logging.info("Creating text chunks...")
    for i, word in enumerate(words):
        counter += 1
        chunk.append(word)
        if word[-1] in {'.', '!', '?'} and counter >= 200:
            text_chunks.append(' '.join(chunk))
            logging.info(f"Chunk #{len(text_chunks)}, chunk length: {counter}")
            # Add overlap by re-adding the last 20 words to the next chunk
            chunk = chunk[-20:] if len(chunk) > 20 else []
            counter = len(chunk)

    if chunk:
        text_chunks.append(' '.join(chunk))
        logging.info(f"Chunk #{len(text_chunks)}, chunk length: {len(chunk)}")

    processed_data = {
        "name": name,
        "chunks": text_chunks,
        "embeddings": []
    }

    logging.info(f"Created {len(text_chunks)} text chunks.")
    return processed_data


async def download_models(models, model_storage):
    try:
        logging.info("Downloading language model...")
        lm_path = f"/app/api/talk_to_docs/models/{models['language_model']}"
        # em_path = f"/app/api/docsAI/models/{models['embedding_model']}"

        service_client = BlobServiceClient(
            account_url=f"https://{model_storage['model_storage_account']}.blob.core.windows.net",
            credential=AzureSasCredential(model_storage['model_sas_token'])
        )

        lm_blob_client = service_client.get_blob_client(container=model_storage['model_container_name'], blob=models['language_model'])
        # em_blob_client = service_client.get_blob_client(container=model_storage['model_container_name'], blob=models['embedding_model'])

        # Download the model files
        with open(lm_path, "wb") as model_download_file:
            model_download_file.write(lm_blob_client.download_blob().readall())

        # with open(em_path, "wb") as model_download_file:
        #     model_download_file.write(em_blob_client.download_blob().readall())

        logging.info(f"Model {models['language_model']} downloaded successfully to {lm_path}.")
        return True
    
    except Exception as e:
        logging.info(f"An error occured while attempting to download the language model: {e}")
        return False


async def embed_query(query, model):
    try:
        logging.info("Embedding query...")
        output = embed.text(
            texts=[query],
            inference_mode='local',
            model=model,
            task_type='search_query',
        )
        embeddings = output['embeddings'][0]
        return embeddings
    except Exception as e:
        logging.error(f"Failed to generate query embeddings: {e}")
        return None
    

async def embed_document(data, model):
    try:
        embeddings = []
        logging.info("Creating embeddings for text chunks...")
        for i, chunk in enumerate(data):
            output = embed.text(
                texts=[chunk],
                inference_mode='local',
                model=model,
                task_type='search_document'
                )
            embedding = output['embeddings'][0]
            logging.info(f"Embedding Length: {len(embedding)}")
            embeddings.append(embedding)
            logging.info(f"Created embedding for chunk {i+1}/{len(data)}")
        return embeddings
    except Exception as e:
        logging.error(f"Failed to generate document embeddings: {e}")
        return None


async def initialize_vector_db(models):
    try:
        # Get the dimentionality settings of embedding model
        dim = len(await embed_query("sample", models['embedding_model']))
        return FAISSConnector(dim)
    except Exception as e:
        logging.error(f"Failed to initiate vector database: {e}")
        return None
    

async def download_and_process_docs(doc_storage, models, db):
    try:
        service_client = BlobServiceClient(
            account_url=f"https://{doc_storage['docs_storage_account']}.blob.core.windows.net",
            credential=AzureSasCredential(doc_storage['docs_sas_token'])
        )
        container_client = service_client.get_container_client(doc_storage['docs_container_name'])
        blob_names = container_client.list_blob_names()

        n_docs = 0
        for blob in blob_names:
            logging.info(f"Downloading blob: {blob}...")
            blob_client = service_client.get_blob_client(container=doc_storage['docs_container_name'], blob=blob)
            blob_data = blob_client.download_blob().readall()

            suffix = os.path.splitext(blob)[1]
            if suffix not in {'.pdf', '.txt', '.docx', '.odt', '.csv', '.xlsx'}:
                logging.info(f"File {blob} is an incompatible format.")
                continue

            processed_doc = await process_document(BytesIO(blob_data), suffix, blob)
            embeddings = await embed_document(processed_doc['chunks'], models['embedding_model'])
            processed_doc['embeddings'] = embeddings

            await db.insert_data(processed_doc)
            n_docs += 1
        logging.info(f"Successfully downloaded and embedded {n_docs} documents.")
        return True
    except Exception as e:
        logging.error(f"Failed to download blobs: {e}")
        return False


async def query_llm(model, model_settings, query, references):
    try:
        logging.info("Prompting language model...")
        model_path = '/app/api/talk_to_docs/models/'

        refs = f"REFERENCES: {references}"
        q = f"QUERY: {query}"
        message = f"{refs} \n\n {q}"

        #TODO: Add n_threads
        model = GPT4All(model_name=model, model_path=model_path, allow_download=False, n_ctx=model_settings['ctx_window'])

        with model.chat_session(system_prompt=model_settings['lm_context']):
            res = model.generate(prompt=message, max_tokens=model_settings['max_tokens'], temp=model_settings['temp'], top_k=model_settings['top_k'], top_p=model_settings['top_p'],
                                            min_p=model_settings['min_p'], repeat_penalty=model_settings['repeat_penalty'], repeat_last_n=model_settings['repeat_last_n'],
                                            n_batch=model_settings['n_batch'], n_predict=model_settings['n_predict'], streaming=model_settings['streaming'])
        return res
    except Exception as e:
        logging.error(f"An error occured while prompting the language model: {e}")
        return None


async def handle_upload():
    logging.info("UPLOAD HANDLER")
    return

async def handle_query(channel_id, query, db, envs):
    try:
        logging.info(f"Incoming query: channel_id: {channel_id} message: {query}")
        query_embedding = await embed_query(query, envs['models']['embedding_model'])
        _, indices = await db.search(query_embedding, 2)

        ref_chunks = ''
        n = 1
        for idx_list in indices:
            for idx in idx_list:
                ref_chunks += f"#{n} - {db.data_dict['chunks'][idx]} \n\n"
                n += 1

        res = await query_llm(envs['models']['language_model'], envs['model_settings'], query, ref_chunks)

        logging.info(f"Language Model Response: {res}")
        return {"answer": res,
                "answer_type": "text",
                "unaccounted": None,
                "addition_buttons": None,
                "buttons": None,
                "images": None,
                "card_data": None}
    except Exception as e:
        return {"answer": e,
                "answer_type": "text",
                "unaccounted": None,
                "addition_buttons": None,
                "buttons": None,
                "images": None,
                "card_data": None}


async def initialize():
    try:
        logging.info("Initializing...")
        envs = await get_envs()
        db = await initialize_vector_db(envs['models'])
        await download_models(envs['models'], envs['models_storage'])
        await download_and_process_docs(envs['doc_storage'], envs['models'], db)
        logging.info("Initialization successful")
        return db, envs
    except Exception as e:
        logging.error(f"Initialization failed: {e}")

# def test():
#     try:
#         global db
#         envs = get_envs()
#         db = initialize_vector_db(envs['models'])
#         #download_models(envs['models'], envs['models_storage'])
#         download_and_process_docs(envs['doc_storage'], envs['models'])
        
#         query = "What is nlsql?"
#         embedded_q = embed_query(query, envs['models']['embedding_model'])
#         _, indices = db.search(embedded_q, 2)

#         chunks = ''
#         n = 1
#         for idx_list in indices:
#             for idx in idx_list:
#                 chunks += f"#{n} - {db.data_dict['chunks'][idx]} \n\n"
#                 n += 1
        
#         res = query_llm(envs['models']['language_model'], envs['model_settings'], query, chunks)

#         logging.info(f"LLM RESPONSE: {res}")
#     except Exception as e:
#         print(e)

